"""Convierte los Excel de ciclos de `Geografia/` en la ficha de cada ciclo.

Paso de UNA sola vez, como los otros dos `preparar_*`: su salida
(`config/ciclos.json`) queda dentro del proyecto.

`historico_electrohuila.ciclo` viene poblado al 100 % (96 valores distintos,
rango 1–200) y es la única columna de la tabla que baja del municipio: describe
cómo está sectorizada la operación de suspensión. De los dos Excel sale su
ficha:

* **`Ciclos de Suspensión.xlsx`** — 95 ciclos con su zona operativa, el número
  de **usuarios** y una descripción del área que cubren. Los usuarios son lo más
  valioso: son el denominador que convierte un conteo de órdenes en una tasa.
* **`SECTORES-2.xlsx`** — qué municipios toca cada ciclo (72 de los 96) y los
  nombres de barrio o sector de cada uno.

Un aviso sobre `SECTORES-2`: está incompleto y desactualizado. Contra los datos
reales, el 87 % de los pares (ciclo, municipio) coinciden y un 1,5 % no, pero
esas discrepancias se concentran en los ciclos rurales grandes (60 y 80), que en
la base tocan más municipios de los que el Excel lista. Por eso la relación
ciclo→municipio se guarda como **referencia**, nunca para descartar una orden.
"""
from __future__ import annotations

import collections
import json
import os
import sys
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "config" / "ciclos.json"
GEOGRAFIA_POR_DEFECTO = RAIZ.parent / "Electro huila" / "Geografia" / "Helectro Huila"

ALIAS = {"ELPITAL": "PITAL", "BELALCAZARPAEZ": "PAEZ", "BELALCAZAR": "PAEZ"}


def clave(texto: object) -> str:
    if texto is None:
        return ""
    s = unicodedata.normalize("NFD", str(texto))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    k = "".join(c for c in s.upper() if c.isalnum())
    return ALIAS.get(k, k)


def entero(v: object) -> int | None:
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return None


def main() -> int:
    origen = Path(os.environ.get("EH_GEOGRAFIA_DIR", "").strip() or GEOGRAFIA_POR_DEFECTO)
    if not origen.exists():
        print(f"No encuentro la carpeta de insumos: {origen}", file=sys.stderr)
        return 2
    print(f"Insumos: {origen}")

    # --- Ficha de cada ciclo ---------------------------------------------
    wb = openpyxl.load_workbook(origen / "Ciclos de Suspensión.xlsx",
                                read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    filas = list(ws.iter_rows(values_only=True))
    cab = [str(c).strip() if c else "" for c in filas[0]]
    fichas: dict[str, dict] = {}
    for fila in filas[1:]:
        d = dict(zip(cab, fila))
        ci = entero(d.get("Ciclo"))
        if ci is None:
            continue
        zona = str(d.get("Zona") or "").strip()
        fichas[str(ci)] = {
            # "1-NORTE" -> "Norte"
            "zona": zona.split("-", 1)[-1].title() if "-" in zona else zona.title(),
            "usuarios": entero(d.get("Usuarios")) or 0,
            "descripcion": str(d.get("Municipios") or "").strip(),
            "ubicacion": str(d.get("Ubicación") or "").strip().upper(),
            "municipios": [],
        }
    wb.close()
    print(f"  Ciclos de Suspensión: {len(fichas)} ciclos · "
          f"{sum(f['usuarios'] for f in fichas.values()):,} usuarios")

    # --- Municipios que toca cada ciclo (referencia) ----------------------
    wb = openpyxl.load_workbook(origen / "SECTORES-2.xlsx", read_only=True, data_only=True)
    filas = list(wb["TODO"].iter_rows(values_only=True))
    cab = [str(c).strip() if c else f"col{i}" for i, c in enumerate(filas[0])]
    munis: dict[int, set[str]] = collections.defaultdict(set)
    for fila in filas[1:]:
        d = dict(zip(cab, fila))
        ci = entero(d.get("CICLO"))
        k = clave(d.get("NOMBRE MPIO"))
        if ci is not None and k and k != "CAUCA":
            munis[ci].add(k)
    wb.close()

    for ci, ms in munis.items():
        fichas.setdefault(str(ci), {
            "zona": "", "usuarios": 0, "descripcion": "", "ubicacion": "", "municipios": [],
        })
        fichas[str(ci)]["municipios"] = sorted(ms)
    confinados = sum(1 for f in fichas.values() if len(f["municipios"]) == 1)
    print(f"  SECTORES-2: {len(munis)} ciclos con municipios · "
          f"{confinados} confinados a uno solo")

    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    with open(SALIDA, "w", encoding="utf-8") as fh:
        json.dump({
            "origen": ["Ciclos de Suspensión.xlsx", "SECTORES-2.xlsx"],
            "ciclos": fichas,
        }, fh, ensure_ascii=False, indent=1)
    print(f"-> {SALIDA.relative_to(RAIZ)} ({SALIDA.stat().st_size / 1024:.1f} KB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
