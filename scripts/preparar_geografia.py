"""Convierte los archivos crudos de `Geografia/` en las capas que consume el mapa.

Es un paso de UNA sola vez: su salida (`web/public/geojson/`) queda dentro del
proyecto, así que el visor y el ETL no vuelven a necesitar la carpeta `Geografia`
original. Se vuelve a correr solo si cambian los límites oficiales.

Entradas (carpeta `Geografia/Helectro Huila`, configurable con EH_GEOGRAFIA_DIR):

    raw/municipios_electrohuila.geojson   42 municipios del área de ElectroHuila
                                          (37 del Huila + Alpujarra, Ataco,
                                          Planadas, Páez y San Vicente del Caguán)
    raw/huila_perimetro_raw.geojson       85 perímetros urbanos / centros poblados
    Municipios.xlsx                       municipio -> zona operativa

Salidas (`web/public/geojson/`):

    municipios.geojson   {nombre, codigo, departamento, zona, lat, lon}
    zonas.geojson        {zona, color}  -- disolución de municipios por zona
    perimetros.geojson   {nombre, municipio, codigo_municipio}

Las zonas replican `Geografia/Helectro Huila/geo_zonas_ajuste.py`: parte del
Excel y luego aplica los ajustes que ahí quedaron documentados (Guadalupe,
Altamira y Gigante completos a Centro; Suaza y Tarqui partidos por la mitad
norte/sur; Pital partido oriente/occidente). Se reimplementa con shapely en vez
de geopandas para no arrastrar esa dependencia.
"""
from __future__ import annotations

import json
import os
import sys
import unicodedata
from pathlib import Path

import openpyxl
from shapely.geometry import box, mapping, shape
from shapely.ops import unary_union

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = RAIZ / "web" / "public" / "geojson"

# Por defecto se busca la carpeta de insumos como hermana del proyecto: los
# archivos crudos pesan ~30 MB y son material de referencia, no del proyecto.
GEOGRAFIA_POR_DEFECTO = RAIZ.parent / "Electro huila" / "Geografia" / "Helectro Huila"

# Tolerancia de simplificación en grados. 0.0005° ≈ 55 m: invisible a la escala
# a la que se mira el mapa y baja el peso del payload en un orden de magnitud.
TOL_MUNICIPIO = 0.0005
TOL_ZONA = 0.0008
TOL_PERIMETRO = 0.0001

# Color de cada zona operativa en el mapa (paleta del visor).
COLOR_ZONA = {
    "Norte": "#4A9EE8",
    "Centro": "#2BD98C",
    "Sur": "#F0C040",
    "Occidente": "#B57BE8",
}

# Municipios que el Excel no lista y que el ajuste de zonas asigna a Centro.
FORZAR_CENTRO = ("Guadalupe", "Altamira", "Gigante")

# Municipios partidos por la mitad (igual área): eje, zona del lado bajo, zona
# del lado alto. Eje "y": bajo = sur. Eje "x": bajo = occidente.
PARTIDOS = {
    "SUAZA": ("y", "Sur", "Centro"),
    "TARQUI": ("y", "Sur", "Centro"),
    "PITAL": ("x", "Occidente", "Centro"),
}

# El Excel escribe algunos municipios con otro nombre que el IGAC.
ALIAS_ZONA = {
    "ELPITAL": "PITAL",
    "BELALCAZARPAEZ": "PAEZ",
    "SALADOBLANCO": "SALADOBLANCO",
    "SANTAMARIA": "SANTAMARIA",
}


def clave(texto: object) -> str:
    """Nombre de municipio reducido a letras y números sin tildes."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFD", str(texto))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return "".join(c for c in s.upper() if c.isalnum())


def kmap(texto: object) -> str:
    k = clave(texto)
    return ALIAS_ZONA.get(k, k)


def leer_geojson(path: Path) -> dict:
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def redondear(obj, nd: int = 5):
    """Redondea todas las coordenadas de una geometría GeoJSON."""
    if isinstance(obj, (list, tuple)):
        if obj and isinstance(obj[0], (int, float)):
            return [round(float(obj[0]), nd), round(float(obj[1]), nd)]
        return [redondear(x, nd) for x in obj]
    return obj


def feature(geom, props: dict) -> dict:
    return {
        "type": "Feature",
        "properties": props,
        "geometry": {"type": geom["type"], "coordinates": redondear(geom["coordinates"])},
    }


def escribir(path: Path, features: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        json.dump({"type": "FeatureCollection", "features": features}, fh,
                  separators=(",", ":"), ensure_ascii=False)
    print(f"  -> {path.name}: {len(features)} polígonos, {path.stat().st_size / 1e6:.2f} MB")


def partir_por_mitad(geom, eje: str):
    """Parte una geometría en dos mitades de igual área (búsqueda binaria)."""
    minx, miny, maxx, maxy = geom.bounds
    objetivo = geom.area / 2
    lo, hi = (miny, maxy) if eje == "y" else (minx, maxx)
    for _ in range(45):
        mid = (lo + hi) / 2
        caja = box(minx, miny, maxx, mid) if eje == "y" else box(minx, miny, mid, maxy)
        if geom.intersection(caja).area < objetivo:
            lo = mid
        else:
            hi = mid
    mid = (lo + hi) / 2
    if eje == "y":
        return geom.intersection(box(minx, miny, maxx, mid)), geom.intersection(box(minx, mid, maxx, maxy))
    return geom.intersection(box(minx, miny, mid, maxy)), geom.intersection(box(mid, miny, maxx, maxy))


def zonas_del_excel(xlsx: Path) -> dict[str, str]:
    """{clave de municipio -> zona} leído de Municipios.xlsx."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    lut: dict[str, str] = {}
    for muni, zona in ws.iter_rows(min_row=2, max_col=2, values_only=True):
        k = kmap(muni)
        if not k or k in ("MUNICIPIO", "CAUCA"):
            continue
        z = str(zona).strip().title() if zona else ""
        if z:
            lut[k] = z
    wb.close()
    for nombre in FORZAR_CENTRO:
        lut[kmap(nombre)] = "Centro"
    return lut


def main() -> int:
    origen = Path(os.environ.get("EH_GEOGRAFIA_DIR", "").strip() or GEOGRAFIA_POR_DEFECTO)
    if not origen.exists():
        print(f"No encuentro la carpeta de insumos: {origen}\n"
              f"Defínela con la variable EH_GEOGRAFIA_DIR.", file=sys.stderr)
        return 2

    print(f"Insumos: {origen}")
    gj_muni = leer_geojson(origen / "raw" / "municipios_electrohuila.geojson")
    lut = zonas_del_excel(origen / "Municipios.xlsx")

    # --- municipios -------------------------------------------------------
    municipios: list[dict] = []
    piezas_zona: dict[str, list] = {}
    sin_zona: list[str] = []

    for f in gj_muni["features"]:
        p = f["properties"]
        nombre = p["municipio_nombre"]
        k = kmap(nombre)
        g = shape(f["geometry"]).buffer(0)

        if k in PARTIDOS:
            eje, z_bajo, z_alto = PARTIDOS[k]
            bajo, alto = partir_por_mitad(g, eje)
            piezas_zona.setdefault(z_bajo, []).append(bajo.buffer(0))
            piezas_zona.setdefault(z_alto, []).append(alto.buffer(0))
            # La zona que se le atribuye al municipio como unidad es la de la
            # mitad más grande; el mapa de zonas sí dibuja las dos.
            zona = z_bajo if bajo.area >= alto.area else z_alto
        elif k in lut:
            zona = lut[k]
            piezas_zona.setdefault(zona, []).append(g)
        else:
            zona = ""
            sin_zona.append(nombre)

        punto = g.representative_point()
        simple = g.simplify(TOL_MUNICIPIO, preserve_topology=True).buffer(0)
        if simple.is_empty:
            simple = g
        municipios.append(feature(mapping(simple), {
            "nombre": nombre,
            "codigo": p.get("municipio_codigo"),
            "departamento": p.get("departamento_nombre"),
            "zona": zona,
            "lat": round(punto.y, 5),
            "lon": round(punto.x, 5),
        }))

    if sin_zona:
        print(f"  AVISO · municipios sin zona en el Excel: {sorted(sin_zona)}")

    escribir(SALIDA / "municipios.geojson", municipios)

    # --- zonas ------------------------------------------------------------
    zonas: list[dict] = []
    for zona in sorted(piezas_zona):
        union = unary_union(piezas_zona[zona]).buffer(0)
        simple = union.simplify(TOL_ZONA, preserve_topology=True).buffer(0)
        if simple.is_empty:
            simple = union
        zonas.append(feature(mapping(simple), {
            "zona": zona,
            "color": COLOR_ZONA.get(zona, "#97999B"),
        }))
    escribir(SALIDA / "zonas.geojson", zonas)
    print(f"     zonas: {[z['properties']['zona'] for z in zonas]}")

    # --- perímetros urbanos / centros poblados ----------------------------
    por_codigo = {}
    for f in gj_muni["features"]:
        cod = str(f["properties"].get("municipio_codigo") or "")
        por_codigo[cod] = f["properties"]["municipio_nombre"]

    gj_per = leer_geojson(origen / "raw" / "huila_perimetro_raw.geojson")
    perimetros: list[dict] = []
    for f in gj_per["features"]:
        p = f["properties"]
        cod = str(p.get("codigo_municipio") or "")
        muni = por_codigo.get(cod, "")
        nombre = (p.get("NOMBRE_GEOGRAFICO") or muni or "Sin nombre").strip()
        g = shape(f["geometry"]).buffer(0)
        simple = g.simplify(TOL_PERIMETRO, preserve_topology=True).buffer(0)
        if simple.is_empty:
            simple = g
        perimetros.append(feature(mapping(simple), {
            "nombre": nombre.upper(),
            "municipio": muni,
            "codigo_municipio": cod,
        }))
    escribir(SALIDA / "perimetros.geojson", perimetros)

    print("Listo.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
